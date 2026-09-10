"""300行×30列規模のExcel表データフィクスチャを生成するスクリプト。

evals/cases/excel-skills/ のケースが参照する inventory_report.xlsx を生成する。
excel-readスキルの`--columns`（2026-09-10追加、列絞り込み）と`--offset`/
`--limit`（既定200行）を組み合わせて、列数・行数どちらも多い表を正しく
分割して読めるか、また読んだ内容から意図的に仕込んだ異常値・不整合を
発見できるかを評価する目的のフィクスチャ。

30列構成（ID/商品名/カテゴリ/単価 + 月別数量12列 + 年間数量合計 +
月別金額12列 + 年間金額合計）、300データ行（2〜301行目、1行目はヘッダー）。
既定の`--limit`200行では1回のread_excel呼び出しで全行を読み切れない。

意図的に仕込んだ異常（すべて決定論的、seed固定）:
1. 121行目: Q列（年間数量合計）の数式がP列（12月）を含まず`=SUM(E121:O121)`
   （本来`=SUM(E121:P121)`のはず）。実データの合計値と数式の計算結果が
   食い違う。
2. 81行目: K列（7月数量）がマイナス値（-120）。数量としてありえない。
3. 51行目と251行目: ID列が同じ値"P-1050"（本来IDは一意のはずの重複）。
4. 151行目: AD列（年間金額合計）の数式が自身を含む`=SUM(R151:AD151)`
   （循環参照。read_excel.pyのwarningsが自動検出する項目）。
5. 201行目: D列（単価）が0なのに、月別金額列（R〜AC）には非ゼロの値が
   入っている（単価0なら売上金額も0のはずという矛盾）。

異常を配置した行（51,81,121,151,201,251）は、既定`--limit`200行での
分割読み込み（1回目1-200行、2回目201-300行）どちらの範囲にも異常が
含まれるよう分散させてある。

使い方:
    python evals/fixtures/generate_excel_columns_scale_fixture.py
"""

from __future__ import annotations

import random
from pathlib import Path

from openpyxl import Workbook
from openpyxl.utils import get_column_letter

FIXTURE_ROOT = Path(__file__).resolve().parent / "excel_columns_scale"

_CATEGORIES = ["食品", "日用品", "衣料品", "家電", "文具"]
_TOTAL_ROWS = 300  # ヘッダー除くデータ行数
_MONTH_QTY_START_COL = 5   # E列（1月数量）
_MONTH_AMOUNT_START_COL = 18  # R列（1月金額）
_QTY_TOTAL_COL = 17  # Q列
_AMOUNT_TOTAL_COL = 30  # AD列

# 異常を仕込むデータ行番号（1始まり、ヘッダーを除く）
_ANOMALY_SUM_RANGE_SHORT_ROW = 120   # 異常1: Q列の数式がP列を含まない
_ANOMALY_NEGATIVE_QTY_ROW = 80       # 異常2: 7月数量がマイナス
_ANOMALY_DUPLICATE_ID_ROWS = (50, 250)  # 異常3: ID重複
_ANOMALY_CIRCULAR_REF_ROW = 150      # 異常4: AD列が自身を含む循環参照
_ANOMALY_ZERO_PRICE_ROW = 200        # 異常5: 単価0なのに金額が非ゼロ


def _header_row() -> list[str]:
    header = ["ID", "商品名", "カテゴリ", "単価"]
    header += [f"数量_{m}月" for m in range(1, 13)]
    header += ["年間数量合計"]
    header += [f"金額_{m}月" for m in range(1, 13)]
    header += ["年間金額合計"]
    return header


def build_fixture(out_dir: Path = FIXTURE_ROOT) -> Path:
    out_dir.mkdir(parents=True, exist_ok=True)
    rng = random.Random(20260910)

    wb = Workbook()
    ws = wb.active
    ws.title = "在庫売上一覧"
    ws.append(_header_row())
    # 既定列幅（8.43）のままだと数式文字列（例"=SUM(E2:P2)"）の長さが列幅を
    # 超え、read_excel.pyの列幅超過警告が全行に大量発生して本質的な異常
    # （循環参照等）の警告が埋もれてしまうため、広めに固定する。
    for col in range(1, 31):
        ws.column_dimensions[get_column_letter(col)].width = 20

    for data_row in range(1, _TOTAL_ROWS + 1):
        excel_row = data_row + 1
        product_id = f"P-{1000 + data_row}"
        if data_row == _ANOMALY_DUPLICATE_ID_ROWS[1]:
            product_id = f"P-{1000 + _ANOMALY_DUPLICATE_ID_ROWS[0]}"  # 異常3: 前半と同じIDを再利用

        category = _CATEGORIES[data_row % len(_CATEGORIES)]
        unit_price = rng.randint(100, 5000)
        if data_row == _ANOMALY_ZERO_PRICE_ROW:
            unit_price = 0  # 異常5: 単価0

        ws.cell(row=excel_row, column=1, value=product_id)
        ws.cell(row=excel_row, column=2, value=f"商品{data_row:03d}")
        ws.cell(row=excel_row, column=3, value=category)
        ws.cell(row=excel_row, column=4, value=unit_price)

        month_qtys = []
        for m in range(12):
            qty = rng.randint(10, 200)
            if data_row == _ANOMALY_NEGATIVE_QTY_ROW and m == 6:  # 7月（0始まり6番目）
                qty = -120  # 異常2: マイナス数量
            month_qtys.append(qty)
            ws.cell(row=excel_row, column=_MONTH_QTY_START_COL + m, value=qty)

        qty_first_letter = get_column_letter(_MONTH_QTY_START_COL)
        qty_last_col = _MONTH_QTY_START_COL + 11
        qty_last_letter = get_column_letter(qty_last_col)
        if data_row == _ANOMALY_SUM_RANGE_SHORT_ROW:
            # 異常1: 12月（P列）を含まず11月（O列）までしか合算しない
            qty_last_letter_short = get_column_letter(qty_last_col - 1)
            ws.cell(
                row=excel_row, column=_QTY_TOTAL_COL,
                value=f"=SUM({qty_first_letter}{excel_row}:{qty_last_letter_short}{excel_row})",
            )
        else:
            ws.cell(
                row=excel_row, column=_QTY_TOTAL_COL,
                value=f"=SUM({qty_first_letter}{excel_row}:{qty_last_letter}{excel_row})",
            )

        for m in range(12):
            # 単価0の異常行でも金額列は通常通り非ゼロの値を入れ、矛盾を作る
            price_for_amount = unit_price if unit_price > 0 else rng.randint(100, 5000)
            amount = month_qtys[m] * price_for_amount if month_qtys[m] > 0 else rng.randint(1000, 50000)
            ws.cell(row=excel_row, column=_MONTH_AMOUNT_START_COL + m, value=amount)

        amount_first_letter = get_column_letter(_MONTH_AMOUNT_START_COL)
        amount_last_col = _MONTH_AMOUNT_START_COL + 11
        amount_last_letter = get_column_letter(amount_last_col)
        amount_total_letter = get_column_letter(_AMOUNT_TOTAL_COL)
        if data_row == _ANOMALY_CIRCULAR_REF_ROW:
            # 異常4: 合計セル自身（AD列）を範囲に含む循環参照
            ws.cell(
                row=excel_row, column=_AMOUNT_TOTAL_COL,
                value=f"=SUM({amount_first_letter}{excel_row}:{amount_total_letter}{excel_row})",
            )
        else:
            ws.cell(
                row=excel_row, column=_AMOUNT_TOTAL_COL,
                value=f"=SUM({amount_first_letter}{excel_row}:{amount_last_letter}{excel_row})",
            )

    out_path = out_dir / "inventory_report.xlsx"
    wb.save(out_path)
    return out_path


def main() -> None:
    path = build_fixture()
    print(f"生成完了: {path}")


if __name__ == "__main__":
    main()
