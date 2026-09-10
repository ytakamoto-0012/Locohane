"""read_excel.pyの`--columns`（列絞り込み）の回帰テスト。

背景（2026-09-10）: 低パラメータのローカルLLMが行を丸ごと取り込む傾向が
強く、列数の多いシート（例300行×30列）を読むと一度に大量のセルが
コンテキストに乗り、情報が抜け落ちた要約をしてしまう問題があった。
対策として読みたい列だけに絞り込める`--columns`を追加した。

単純に列を間引くと「rows配列の何番目が元々何列目だったか」の対応が
崩れるため、トップレベルに列アルファベット配列`columns`を返し、
`rows[i][j]`が常に`columns[j]`列の値になるようにした。この対応が
`--style`時の`column_widths`と`warnings`（列幅超過・循環参照）でも
崩れていないことを確認する（特にwarningsは、間引かれた列でも実列番号を
正しく指す必要がある）。
"""

import sys
from pathlib import Path

import openpyxl
import pytest

_SCRIPTS_DIR = Path(__file__).resolve().parent.parent / "skills" / "excel-read" / "scripts"
if str(_SCRIPTS_DIR) not in sys.path:
    sys.path.insert(0, str(_SCRIPTS_DIR))

_OFFICE_SHARED = Path(__file__).resolve().parent.parent / "skills" / "office_shared"
if str(_OFFICE_SHARED) not in sys.path:
    sys.path.insert(0, str(_OFFICE_SHARED))

from excel_common import parse_column_selection  # noqa: E402
from read_excel import _read_xls, _read_xlsx  # noqa: E402


class TestParseColumnSelection:
    def test_none_arg_returns_none(self):
        assert parse_column_selection(None, total_columns=10) is None

    def test_alphabet_tokens_resolved_to_indices(self):
        assert parse_column_selection("A,C", total_columns=10) == [1, 3]

    def test_numeric_tokens_resolved_to_indices(self):
        assert parse_column_selection("1,3", total_columns=10) == [1, 3]

    def test_unordered_input_sorted_ascending(self):
        assert parse_column_selection("C,A", total_columns=10) == [1, 3]

    def test_duplicate_tokens_deduplicated(self):
        assert parse_column_selection("A,A,C", total_columns=10) == [1, 3]

    def test_index_equal_to_total_columns_is_valid(self):
        assert parse_column_selection("J", total_columns=10) == [10]

    def test_index_beyond_total_columns_raises(self):
        with pytest.raises(ValueError, match="範囲外"):
            parse_column_selection("Z", total_columns=10)

    def test_empty_string_raises(self):
        with pytest.raises(ValueError):
            parse_column_selection(",", total_columns=10)


class TestReadXlsxColumnsFilter:
    def _make_workbook(self, tmp_path):
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Sheet1"
        ws["A1"] = "id"
        ws["B1"] = "name"
        ws["C1"] = "value"
        path = tmp_path / "book.xlsx"
        wb.save(path)
        return path

    def test_columns_key_lists_all_columns_by_default(self, tmp_path):
        path = self._make_workbook(tmp_path)
        result = _read_xlsx(path, sheet_arg="Sheet1", offset=0, limit=200, data_only=False)
        assert result["columns"] == ["A", "B", "C"]
        assert result["rows"][0] == [{"value": "id"}, {"value": "name"}, {"value": "value"}]

    def test_columns_filter_narrows_rows_in_matching_order(self, tmp_path):
        path = self._make_workbook(tmp_path)
        result = _read_xlsx(path, sheet_arg="Sheet1", offset=0, limit=200, data_only=False, columns="C,A")
        assert result["columns"] == ["A", "C"]
        assert result["rows"][0] == [{"value": "id"}, {"value": "value"}]
        # total_columns/total_rowsはシート全体基準のまま変わらない
        assert result["total_columns"] == 3

    def test_columns_without_sheet_raises(self, tmp_path):
        path = self._make_workbook(tmp_path)
        with pytest.raises(ValueError, match="--sheet"):
            _read_xlsx(path, sheet_arg=None, offset=0, limit=200, data_only=False, columns="A")

    def test_column_widths_limited_to_selected_columns(self, tmp_path):
        path = self._make_workbook(tmp_path)
        result = _read_xlsx(path, sheet_arg="Sheet1", offset=0, limit=200, data_only=False, columns="C,A")
        assert set(result["column_widths"].keys()) == {"A", "C"}

    def test_warning_cell_ref_uses_actual_column_after_filtering(self, tmp_path):
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Sheet1"
        ws["B8"] = 1
        ws["N8"] = "=SUM(B8:N8)"
        path = tmp_path / "book.xlsx"
        wb.save(path)

        # N列（実列番号14）だけに絞り込んでも、配列内の位置は0番目になる。
        # 修正前は「配列位置+1」を実列番号として使っていたため、ここが
        # "A8"のような誤った列を指してしまうバグがあった。
        result = _read_xlsx(
            path, sheet_arg="Sheet1", offset=0, limit=200, data_only=False, columns="N",
        )

        assert result["columns"] == ["N"]
        assert "warnings" in result
        assert any("循環参照" in w and "N8" in w for w in result["warnings"])


class TestReadXlsColumnsWithoutSheetRaises:
    def test_columns_without_sheet_raises_before_opening_file(self, tmp_path):
        """_read_xlsのガードは、シート一覧モードの早期returnより前に置く必要がある
        （置き場所を間違えると--sheet省略時に無条件でシート一覧を返してしまい、
        --columns指定が黙って無視される。ガードが正しく先頭にあれば、実在しない
        パスでもxlrd.open_workbookに到達する前にValueErrorが送出されるはず）。
        """
        missing_path = tmp_path / "does_not_exist.xls"
        with pytest.raises(ValueError, match="--sheet"):
            _read_xls(missing_path, sheet_arg=None, offset=0, limit=200, columns="A")
