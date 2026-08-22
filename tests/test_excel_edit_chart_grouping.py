"""excel-editスキル（skills/excel-edit/scripts/_ops.py）のグラフgrouping（積み上げ）対応。

背景: 「積み上げ棒グラフ」を作りたいタスクで、add_chartのtype一覧
（bar/line/pie/scatter）に積み上げを指定する手段が無く、LLMが
openpyxlのソースをGrepして探し回った末に思考ループへ突入した
（2026-08-22 app.log調査）。openpyxlのBarChart自体は`grouping`属性
（clustered/stacked/percentStacked）を持つため、add_chart/update_chartの
opとして薄く公開する。
"""

import sys
from pathlib import Path

import openpyxl
import pytest

_SCRIPTS_DIR = Path(__file__).resolve().parent.parent / "skills" / "excel-edit" / "scripts"
if str(_SCRIPTS_DIR) not in sys.path:
    sys.path.insert(0, str(_SCRIPTS_DIR))

from _ops import apply_op  # noqa: E402


def _make_wb_with_data():
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Sheet1"
    apply_op(
        wb,
        {
            "op": "set_range",
            "sheet": "Sheet1",
            "start_cell": "A1",
            "rows": [["月", "収入", "支出"], ["1月", 100, 50], ["2月", 120, 60]],
        },
    )
    return wb, ws


class TestAddChartGrouping:
    def test_stacked_sets_grouping_and_overlap(self):
        wb, ws = _make_wb_with_data()
        apply_op(
            wb,
            {
                "op": "add_chart",
                "sheet": "Sheet1",
                "type": "bar",
                "data_range": "B1:C3",
                "anchor": "E1",
                "grouping": "stacked",
            },
        )
        chart = ws._charts[0]
        assert chart.grouping == "stacked"
        assert chart.overlap == 100

    def test_percent_stacked_sets_grouping_and_overlap(self):
        wb, ws = _make_wb_with_data()
        apply_op(
            wb,
            {
                "op": "add_chart",
                "sheet": "Sheet1",
                "type": "bar",
                "data_range": "B1:C3",
                "anchor": "E1",
                "grouping": "percentStacked",
            },
        )
        chart = ws._charts[0]
        assert chart.grouping == "percentStacked"
        assert chart.overlap == 100

    def test_clustered_does_not_force_overlap(self):
        wb, ws = _make_wb_with_data()
        apply_op(
            wb,
            {
                "op": "add_chart",
                "sheet": "Sheet1",
                "type": "bar",
                "data_range": "B1:C3",
                "anchor": "E1",
                "grouping": "clustered",
            },
        )
        chart = ws._charts[0]
        assert chart.grouping == "clustered"
        assert chart.overlap is None

    def test_omitted_grouping_leaves_openpyxl_default(self):
        wb, ws = _make_wb_with_data()
        apply_op(
            wb,
            {"op": "add_chart", "sheet": "Sheet1", "type": "bar", "data_range": "B1:C3", "anchor": "E1"},
        )
        chart = ws._charts[0]
        assert chart.grouping == "clustered"

    def test_grouping_on_non_bar_chart_raises(self):
        wb, ws = _make_wb_with_data()
        with pytest.raises(ValueError, match="type=bar"):
            apply_op(
                wb,
                {
                    "op": "add_chart",
                    "sheet": "Sheet1",
                    "type": "line",
                    "data_range": "B1:C3",
                    "anchor": "E1",
                    "grouping": "stacked",
                },
            )

    def test_invalid_grouping_value_raises(self):
        wb, ws = _make_wb_with_data()
        with pytest.raises(ValueError, match="未対応のgroupingです"):
            apply_op(
                wb,
                {
                    "op": "add_chart",
                    "sheet": "Sheet1",
                    "type": "bar",
                    "data_range": "B1:C3",
                    "anchor": "E1",
                    "grouping": "sideways",
                },
            )


class TestUpdateChartGrouping:
    def test_update_existing_bar_chart_to_stacked(self):
        wb, ws = _make_wb_with_data()
        apply_op(
            wb,
            {"op": "add_chart", "sheet": "Sheet1", "type": "bar", "data_range": "B1:C3", "anchor": "E1"},
        )
        apply_op(wb, {"op": "update_chart", "sheet": "Sheet1", "chart_index": 0, "grouping": "stacked"})

        chart = ws._charts[0]
        assert chart.grouping == "stacked"
        assert chart.overlap == 100

    def test_update_chart_grouping_alone_satisfies_required_arg_check(self):
        wb, ws = _make_wb_with_data()
        apply_op(
            wb,
            {"op": "add_chart", "sheet": "Sheet1", "type": "bar", "data_range": "B1:C3", "anchor": "E1"},
        )
        # grouping単体の指定でも「いずれか1つ以上」の必須チェックを通過できる
        apply_op(wb, {"op": "update_chart", "sheet": "Sheet1", "chart_index": 0, "grouping": "percentStacked"})
        assert ws._charts[0].grouping == "percentStacked"
