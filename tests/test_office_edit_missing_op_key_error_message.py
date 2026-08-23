"""ops適用時にKeyError（必須キー未指定）が発生した場合のエラーメッセージ回帰テスト。

背景（2026-08-23）: excel-vbaマクロブック作成タスクで、workerが`set_cell`opに
`cell`ではなく`row`/`col`キーを渡し、`ops[3]（op='set_cell'）の適用に失敗しました:
'cell'`という素のKeyErrorメッセージだけが返った。「'cell'」だけでは
「cellキーが無い」のか「cellキーの値がおかしい」のか区別できず、LLMが
約30秒間、無関係な原因（シート名の不一致等）を疑って迷走した末に
`read_excel.py`で現状確認するという遠回りをした。

`edit_excel.py`/`edit_vba.py`/`edit_docx.py`はいずれも同じ
`except (KeyError, ValueError, TypeError) as e: f"...の適用に失敗しました: {e}"`
という実装だったため、KeyErrorだけ「必須キー{e}が指定されていません」という
明確なメッセージに変えた（ValueError/TypeErrorは従来通り）。

excel-edit・docx-edit はどちらもスクリプトディレクトリ内に同名の`_ops.py`を
持つため、両方を同じpytestプロセス内でimportすると`sys.modules`のキャッシュに
より後からimportした側が先にimport済みの別ディレクトリの`_ops`を誤って
再利用してしまう。各テストの冒頭で関連モジュールを`sys.modules`から明示的に
取り除いてから絶対パス指定でimportし直すことで、この汚染を避けている。
"""

import importlib
import json
import sys
from pathlib import Path

_EXCEL_EDIT_SCRIPTS = Path(__file__).resolve().parent.parent / "skills" / "excel-edit" / "scripts"
_EXCEL_VBA_EDIT_SCRIPTS = Path(__file__).resolve().parent.parent / "skills" / "excel-vba-edit" / "scripts"
_DOCX_EDIT_SCRIPTS = Path(__file__).resolve().parent.parent / "skills" / "docx-edit" / "scripts"
_OFFICE_SHARED = Path(__file__).resolve().parent.parent / "skills" / "office_shared"


def _fresh_import(module_name: str, scripts_dir: Path):
    """他スキルの同名モジュール（例: _ops）とのsys.modulesキャッシュ衝突を避け、
    指定ディレクトリのモジュールを確実に読み直す。"""
    sys.modules.pop(module_name, None)
    for d in (scripts_dir, _OFFICE_SHARED):
        path_str = str(d)
        if path_str in sys.path:
            sys.path.remove(path_str)
        sys.path.insert(0, path_str)
    return importlib.import_module(module_name)


def test_edit_excel_missing_required_key_gives_clear_message(tmp_path, monkeypatch, capsys):
    sys.modules.pop("_ops", None)
    edit_excel = _fresh_import("edit_excel", _EXCEL_EDIT_SCRIPTS)

    target = tmp_path / "book.xlsx"
    ops = [
        {"op": "add_sheet", "name": "Sheet1"},
        {"op": "set_cell", "sheet": "Sheet1", "row": 1, "col": 1, "value": "x"},  # "cell"キー欠落
    ]
    monkeypatch.setattr(sys, "argv", ["edit_excel.py", str(target), "--new", "--ops-json", json.dumps(ops)])

    exit_code = edit_excel.main()

    assert exit_code == 1
    err = capsys.readouterr().err
    assert "必須キー" in err
    assert "'cell'" in err
    assert "の適用に失敗しました" not in err


def test_edit_vba_missing_required_key_gives_clear_message(tmp_path, monkeypatch):
    sys.modules.pop("_vba_ops", None)
    edit_vba = _fresh_import("edit_vba", _EXCEL_VBA_EDIT_SCRIPTS)

    class _FakeVBProject:
        pass

    class _FakeWorkbook:
        VBProject = _FakeVBProject()

        def Close(self, SaveChanges=False):
            pass

    class _FakeExcel:
        Visible = None
        DisplayAlerts = None
        AutomationSecurity = None
        Hwnd = 0

        class Workbooks:
            @staticmethod
            def Add():
                return _FakeWorkbook()

        def Quit(self):
            pass

    class _FakeWin32ClientModule:
        def DispatchEx(self, prog_id):
            return _FakeExcel()

    class _FakePythoncomModule:
        def CoInitialize(self):
            pass

        def CoUninitialize(self):
            pass

    monkeypatch.setitem(sys.modules, "pythoncom", _FakePythoncomModule())
    monkeypatch.setitem(sys.modules, "win32com.client", _FakeWin32ClientModule())

    target = tmp_path / "book.xlsm"
    # add_moduleはnameが必須（nameを欠落させる）
    ops = [{"op": "add_module", "code": "Sub Foo()\nEnd Sub"}]

    try:
        edit_vba._edit_vba(target, target, ops, is_new=True, overwrite=False)
        raised = None
    except ValueError as e:
        raised = str(e)

    assert raised is not None
    assert "必須キー" in raised
    assert "'name'" in raised
    assert "の適用に失敗しました" not in raised


def test_edit_docx_missing_required_key_gives_clear_message(tmp_path, monkeypatch, capsys):
    sys.modules.pop("_ops", None)
    edit_docx = _fresh_import("edit_docx", _DOCX_EDIT_SCRIPTS)

    import docx

    target = tmp_path / "doc.docx"
    doc = docx.Document()
    doc.add_paragraph("hello")
    doc.save(target)

    ops = [{"op": "find_replace", "new_text": "bye"}]  # "old_text"キー欠落
    monkeypatch.setattr(sys, "argv", ["edit_docx.py", str(target), "--ops-json", json.dumps(ops)])

    exit_code = edit_docx.main()

    assert exit_code == 1
    err = capsys.readouterr().err
    assert "必須キー" in err
    assert "'old_text'" in err
    assert "の適用に失敗しました" not in err
